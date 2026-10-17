# sidecar/tests/make_office_fixtures.py
"""
Genera DOCX y PPTX pequeños con imágenes para las pruebas de OCR en documentos de Office. Sin dependencias nuevas:
el DOCX se escribe a mano (zip + XML) y el PPTX con python-pptx (ya viaja con markitdown).

Bloques de `build_docx`:
    ("p", texto)
    ("img", PNG_o_bytes, ancho_px, alto_px, alt, extensión)      # imagen en línea
    ("table", [[celda, ...], ...])                               # celda = str o ("img", ...)
"""

import io
import zipfile
from html import escape
from pathlib import Path

from PIL import Image

from make_fixtures import DIAGRAM_LINES, ES_LINES, text_page

W_NS = ('xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
        'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships" '
        'xmlns:wp="http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing" '
        'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" '
        'xmlns:pic="http://schemas.openxmlformats.org/drawingml/2006/picture"')
CONTENT_TYPES = {"png": "image/png", "jpg": "image/jpeg", "jpeg": "image/jpeg", "emf": "image/x-emf",
                 "wmf": "image/x-wmf", "svg": "image/svg+xml"}
# Cabecera mínima de un EMF: Pillow lo identifica pero no puede decodificarlo fuera de Windows.
FAKE_EMF = (b"\x01\x00\x00\x00" + b"\x6c\x00\x00\x00" + b"\x00" * 32 + b" EMF" + b"\x00" * 64)


def png_bytes(img: Image.Image, fmt: str = "PNG") -> bytes:
    b = io.BytesIO()
    img.save(b, fmt)
    return b.getvalue()


def img(data: bytes, w: int, h: int, alt: str = "", ext: str = "png"):
    return ("img", data, w, h, alt, ext)


def _drawing(rid: str, n: int, w: int, h: int, alt: str) -> str:
    emu_w, emu_h = w * 9525, h * 9525
    return (f'<w:r><w:drawing><wp:inline><wp:extent cx="{emu_w}" cy="{emu_h}"/>'
            f'<wp:docPr id="{n}" name="Imagen {n}" descr="{escape(alt, quote=True)}"/>'
            '<a:graphic><a:graphicData uri="http://schemas.openxmlformats.org/drawingml/2006/picture"><pic:pic>'
            f'<pic:nvPicPr><pic:cNvPr id="{n}" name="i{n}"/><pic:cNvPicPr/></pic:nvPicPr>'
            f'<pic:blipFill><a:blip r:embed="{rid}"/><a:stretch><a:fillRect/></a:stretch></pic:blipFill>'
            f'<pic:spPr><a:xfrm><a:off x="0" y="0"/><a:ext cx="{emu_w}" cy="{emu_h}"/></a:xfrm><a:prstGeom prst="rect"/></pic:spPr>'
            '</pic:pic></a:graphicData></a:graphic></wp:inline></w:drawing></w:r>')


def build_docx(path: Path, blocks: list) -> Path:
    media: list[tuple[str, bytes]] = []   # (nombre dentro de word/media, bytes)
    exts: set[str] = set()

    def image_run(block) -> str:
        _, data, w, h, alt, ext = block
        exts.add(ext)
        name = f"image{len(media) + 1}.{ext}"
        media.append((name, data))
        return _drawing(f"rId{100 + len(media)}", len(media), w, h, alt)

    def para(content: str) -> str:
        return f"<w:p>{content}</w:p>"

    body = []
    for b in blocks:
        if b[0] == "p":
            body.append(para(f"<w:r><w:t>{escape(b[1])}</w:t></w:r>"))
        elif b[0] == "img":
            body.append(para(image_run(b)))
        elif b[0] == "table":
            rows = []
            for row in b[1]:
                cells = []
                for cell in row:
                    inner = para(image_run(cell)) if isinstance(cell, tuple) else para(f"<w:r><w:t>{escape(cell)}</w:t></w:r>")
                    cells.append(f"<w:tc>{inner}</w:tc>")
                rows.append(f"<w:tr>{''.join(cells)}</w:tr>")
            body.append(f"<w:tbl>{''.join(rows)}</w:tbl>")
    doc = f'<?xml version="1.0" encoding="UTF-8" standalone="yes"?><w:document {W_NS}><w:body>{"".join(body)}</w:body></w:document>'
    defaults = "".join(f'<Default Extension="{e}" ContentType="{CONTENT_TYPES[e]}"/>' for e in sorted(exts))
    ct = ('<?xml version="1.0"?><Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
          '<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>'
          f'<Default Extension="xml" ContentType="application/xml"/>{defaults}'
          '<Override PartName="/word/document.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"/></Types>')
    rels = ('<?xml version="1.0"?><Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
            '<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="word/document.xml"/></Relationships>')
    drels = ('<?xml version="1.0"?><Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
             + "".join(f'<Relationship Id="rId{101 + i}" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/image" Target="media/{n}"/>'
                       for i, (n, _) in enumerate(media)) + "</Relationships>")
    with zipfile.ZipFile(path, "w") as z:
        z.writestr("[Content_Types].xml", ct)
        z.writestr("_rels/.rels", rels)
        z.writestr("word/document.xml", doc)
        z.writestr("word/_rels/document.xml.rels", drels)
        for name, data in media:
            z.writestr(f"word/media/{name}", data)
    return path


def make_docx_corpus(out: Path) -> dict[str, Path]:
    out.mkdir(parents=True, exist_ok=True)
    shot = png_bytes(text_page(ES_LINES, size=44, px=(900, 420)))
    diagram = png_bytes(text_page(DIAGRAM_LINES, size=44, px=(900, 400)))
    icon = png_bytes(Image.new("RGB", (48, 48), "navy"))
    logo = png_bytes(Image.new("RGB", (240, 70), "darkgreen"))
    page = png_bytes(text_page(ES_LINES, size=54))                       # página escaneada completa
    page_180 = png_bytes(text_page(ES_LINES, size=54).rotate(180))
    photo = png_bytes(Image.effect_noise((900, 600), 60).convert("RGB"))  # imagen sin texto
    docs = {
        "solo_texto": [("p", "Un documento sin imágenes."), ("p", "Segundo párrafo con ñandú.")],
        "captura_entre_parrafos": [("p", "Antes de la imagen."), img(shot, 600, 280, "Captura del sistema"), ("p", "Después de la imagen.")],
        "solo_imagen": [img(page, 620, 877)],
        "pagina_escaneada_180": [("p", "Escaneo girado:"), img(page_180, 620, 877)],
        "icono_con_alt": [("p", "Texto."), img(icon, 32, 32, "Correo de contacto")],
        "iconos_y_logo": [("p", "Texto."), img(icon, 32, 32), img(logo, 240, 70), img(icon, 32, 32)],
        "misma_imagen_repetida": [("p", "Uno"), img(shot, 600, 280), ("p", "Dos"), img(shot, 600, 280), ("p", "Tres"), img(shot, 600, 280)],
        "imagen_en_tabla": [("p", "Tabla:"), ("table", [["Celda A", img(diagram, 450, 200)], ["Celda C", "Celda D"]])],
        "alt_automatico": [img(shot, 600, 280, "image.png"), ("p", "x"), img(shot, 500, 233, "Imagen 3")],
        "foto_sin_texto": [("p", "Una foto:"), img(photo, 600, 400)],
        "emf_y_captura": [img(FAKE_EMF, 200, 100, "", "emf"), ("p", "Entre."), img(FAKE_EMF, 200, 100, "", "emf"), img(shot, 600, 280)],
        "imagen_corrupta": [("p", "Hola"), img(b"esto no es una imagen", 300, 200), ("p", "Fin")],
    }
    return {name: build_docx(out / f"{name}.docx", blocks) for name, blocks in docs.items()}


# ----------------------------------------------------------------------------------------------------------------
# PPTX
# ----------------------------------------------------------------------------------------------------------------

def patch_zip(path: Path, replacements: dict[str, bytes], text_replace: dict[str, tuple[str, str]] | None = None) -> None:
    """Reescribe el zip cambiando el contenido de partes (por nombre) o aplicando reemplazos de texto."""
    items = []
    with zipfile.ZipFile(path) as z:
        for info in z.infolist():
            data = z.read(info.filename)
            if info.filename in replacements:
                data = replacements[info.filename]
            if text_replace and info.filename in text_replace:
                old, new = text_replace[info.filename]
                data = data.decode("utf-8").replace(old, new).encode("utf-8")
            items.append((info.filename, data))
    with zipfile.ZipFile(path, "w") as z:
        for name, data in items:
            z.writestr(name, data)


def make_pptx_corpus(out: Path) -> dict[str, Path]:
    from pptx import Presentation
    from pptx.util import Inches

    out.mkdir(parents=True, exist_ok=True)

    def stream(data: bytes):
        return io.BytesIO(data)

    shot = png_bytes(text_page(ES_LINES, size=44, px=(900, 420)))
    diagram = png_bytes(text_page(DIAGRAM_LINES, size=44, px=(900, 400)))
    icon = png_bytes(Image.new("RGB", (48, 48), "navy"))
    page = png_bytes(text_page(ES_LINES, size=54, px=(1280, 720)))
    page_180 = png_bytes(text_page(ES_LINES, size=54, px=(1280, 720)).rotate(180))

    def deck(builder) -> Presentation:
        p = Presentation()
        builder(p)
        return p

    def title_slide(p, title):
        s = p.slides.add_slide(p.slide_layouts[5])
        s.shapes.title.text = title
        return s

    def picture(slide_or_group, data, left, top, width, descr=None):
        shapes = slide_or_group.shapes
        pic = shapes.add_picture(stream(data), Inches(left), Inches(top), width=Inches(width))
        if descr is not None:
            pic._element.nvPicPr.cNvPr.set("descr", descr)
        return pic

    def only_text(p):
        s = title_slide(p, "Solo texto")
        s.shapes.add_textbox(Inches(1), Inches(2), Inches(5), Inches(1)).text_frame.text = "Contenido sin imágenes con ñandú"

    def with_picture(p):
        s = title_slide(p, "Diapositiva con figura")
        s.shapes.add_textbox(Inches(1), Inches(1.5), Inches(5), Inches(0.6)).text_frame.text = "Texto antes de la figura"
        picture(s, shot, 1, 2.5, 6, descr="Captura del sistema")
        s.shapes.add_textbox(Inches(1), Inches(6.5), Inches(5), Inches(0.6)).text_frame.text = "Texto después de la figura"

    def icons(p):
        s = title_slide(p, "Iconos")
        for i in range(3):
            picture(s, icon, 8 + 0.4 * i, 0.2, 0.3)
        picture(s, icon, 8, 1.0, 0.3, descr="Correo de contacto")

    def full_page(p, data):
        s = p.slides.add_slide(p.slide_layouts[6])
        s.shapes.add_picture(stream(data), 0, 0, width=p.slide_width, height=p.slide_height)

    def grouped(p):
        s = title_slide(p, "Imagen en un grupo")
        g = s.shapes.add_group_shape()
        picture(g, diagram, 1, 3, 5)

    def placeholder(p):
        s = p.slides.add_slide(p.slide_layouts[8])   # Picture with Caption
        s.shapes.title.text = "Marcador con imagen"
        ph = next(x for x in s.placeholders if x.placeholder_format.type == 18)
        ph.insert_picture(stream(shot))

    def repeated(p):
        for i in range(1, 4):
            s = title_slide(p, f"Diapositiva {i}")
            picture(s, shot, 1, 2, 6)

    def table_and_notes(p):
        s = title_slide(p, "Tabla y notas")
        t = s.shapes.add_table(2, 2, Inches(1), Inches(1.5), Inches(5), Inches(1)).table
        for r, row in enumerate([["Mes", "Total"], ["Enero", "1200"]]):
            for c, v in enumerate(row):
                t.cell(r, c).text = v
        picture(s, shot, 1, 3.5, 5)
        s.notes_slide.notes_text_frame.text = "Notas del orador"

    def three_distinct(p):
        for i in range(3):
            s = title_slide(p, f"Figura {i + 1}")
            picture(s, png_bytes(text_page(ES_LINES, size=40 + i, px=(900, 420))), 1, 2, 6)

    builders = {
        "solo_texto": only_text, "con_figura": with_picture, "iconos": icons,
        "pagina_completa": lambda p: full_page(p, page), "pagina_completa_180": lambda p: full_page(p, page_180),
        "en_grupo": grouped, "marcador": placeholder, "repetida": repeated,
        "tabla_y_notas": table_and_notes, "tres_distintas": three_distinct,
        "imagen_corrupta": lambda p: picture(title_slide(p, "Rota"), shot, 1, 2, 6),
        "tipo_emf": lambda p: picture(title_slide(p, "Vector"), shot, 1, 2, 6),
    }
    files = {}
    for name, builder in builders.items():
        path = out / f"{name}.pptx"
        deck(builder).save(path)
        files[name] = path
    # Variantes parcheadas: bytes de imagen corruptos y tipo de contenido EMF.
    with zipfile.ZipFile(files["imagen_corrupta"]) as z:
        media = next(n for n in z.namelist() if n.startswith("ppt/media/"))
    patch_zip(files["imagen_corrupta"], {media: b"esto no es una imagen"})
    patch_zip(files["tipo_emf"], {}, {"[Content_Types].xml": ('Extension="png" ContentType="image/png"', 'Extension="png" ContentType="image/x-emf"')})
    return files
