# sidecar/ocr/office_images.py
"""
OCR de las imágenes de DOCX y PPTX. Ver el change add-ocr-office-images.

- Sin imágenes: se delega en `convert_document` (salida idéntica a la de siempre, sin cargar el motor).
- DOCX: mammoth entrega cada imagen en orden de documento; su texto ocupa el lugar de la imagen.
- PPTX: cada imagen se sustituye en una copia en memoria por un cuadro de texto en las mismas coordenadas y la copia
  se convierte con markitdown, de modo que el texto sale en el orden de lectura de la diapositiva.
- Se leen solo imágenes de tamaño mostrado >= OFFICE_MIN_AREA, una vez por imagen distinta, y se descarta el texto
  de menos de OFFICE_MIN_TEXT_CHARS alfanuméricos (iconos y pictogramas). El alt text real se conserva.
"""

import hashlib
import re
import sys
import threading
import traceback
import zipfile
from pathlib import Path
from typing import Callable

from converter import convert_document
from ocr.base import OcrEngine
from ocr.diagnostics import ConversionError, PageIssue, log_issue
from ocr.image_converter import ConversionResult
from ocr.image_loader import load_image
from ocr.page_classifier import SCAN_COVERAGE
from ocr.recognize import read_simple, read_upright

OFFICE_MIN_AREA = 0.01               # fracción de la página/diapositiva que debe ocupar la imagen (calibrado, S10)
OFFICE_MIN_TEXT_CHARS = 6            # alfanuméricos mínimos para emitir el texto de una imagen (S7)
FALLBACK_MIN_SIDE_PX = 100           # DOCX sin tamaño mostrado: ambos lados >= esto ...
FALLBACK_MIN_PIXELS = 20_000         # ... y área >= esto
FALLBACK_PAGE_LIKE_PX = 1000         # DOCX sin tamaño mostrado: ambos lados >= esto = imagen "de página"
SKIPPED_TYPES = {"image/x-emf", "image/emf", "image/x-wmf", "image/wmf", "image/svg+xml", "image/x-emz", "image/x-wmz"}

_AUTO_ALT = re.compile(r"^(picture|imagen|image|foto|figura|graphic|gr[aá]fico)\s*\d*$", re.IGNORECASE)
_FILE_ALT = re.compile(r"\.(png|jpe?g|gif|bmp|tiff?|webp|emf|wmf|svg)$", re.IGNORECASE)
_XML_ILLEGAL = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")
_EMU_PER_INCH = 914_400


class _Cancelled(Exception):
    """Señal interna para salir de mammoth cuando se cancela la conversión."""


# ---------------------------------------------------------------------------------------------------------------
# Utilidades comunes (puras, con tests propios)
# ---------------------------------------------------------------------------------------------------------------

def real_alt(alt: str | None, shape_name: str | None = None) -> str:
    """El texto alternativo solo si es una descripción: no el nombre de la forma, ni un archivo ('image.png'),
    ni un nombre automático ('Imagen 3')."""
    alt = " ".join((alt or "").split())
    if not alt or alt.lower() == (shape_name or "").strip().lower() or _FILE_ALT.search(alt) or _AUTO_ALT.match(alt):
        return ""
    return alt.replace("*", "")


def alnum_count(text: str) -> int:
    return len(re.sub(r"[\W_]", "", text))


def image_paragraphs(alt: str, text: str) -> list[str]:
    """Párrafos que sustituyen a la imagen: el alt text en cursiva y luego el texto reconocido."""
    return ([f"*{alt}*"] if alt else []) + [p for p in text.split("\n\n") if p.strip()]


def _read_embedded(data: bytes, name: str, page_like: bool, engine: OcrEngine) -> tuple[str, int]:
    """Texto de una imagen incrustada y líneas descartadas. Sin texto útil -> ("", n)."""
    image, _ = load_image(data, name)
    try:
        text, dropped = (read_upright if page_like else read_simple)(image, engine)
    except ConversionError as e:
        if e.code == "ocr_failed" and e.cause is None:  # ninguna orientación lee texto (p. ej. foto de fondo): no es un fallo
            return "", 0
        raise
    return (text if alnum_count(text) >= OFFICE_MIN_TEXT_CHARS else ""), dropped


def _pixel_size(data: bytes) -> tuple[int, int] | None:
    from io import BytesIO
    from PIL import Image
    try:
        return Image.open(BytesIO(data)).size
    except Exception:
        return None


class _Reader:
    """Estado compartido de una conversión: motor perezoso, deduplicación, avisos y progreso."""

    def __init__(self, name: str, cancel_flag: threading.Event, on_progress: Callable[[int], None],
                 engine: OcrEngine | None, total: int):
        self.name, self.cancel_flag, self.on_progress, self.engine, self.total = name, cancel_flag, on_progress, engine, max(1, total)
        self.warnings: list[PageIssue] = []
        self.seen: set[str] = set()
        self.skipped = 0
        self.dropped = 0
        self.done = 0

    def finish(self) -> list[PageIssue]:
        if self.skipped:
            self.warnings.append(PageIssue("images_skipped", "decode", count=self.skipped))
        if self.dropped:
            self.warnings.append(PageIssue("low_confidence_dropped", "layout", count=self.dropped))
        return self.warnings

    def read(self, data: bytes, content_type: str | None, relevant: bool, page_like: bool, page: int | None) -> str:
        """Texto OCR de una imagen (o "" si se omite, es duplicada, falla o no tiene texto útil)."""
        if self.cancel_flag.is_set():
            raise _Cancelled()
        digest = hashlib.sha1(data).hexdigest()
        text = ""
        if (content_type or "").lower() in SKIPPED_TYPES:
            self.skipped += 1
        elif relevant and digest not in self.seen:
            self.seen.add(digest)
            stage = "decode"
            try:
                if self.engine is None:
                    from ocr.engine import get_engine  # carga perezosa del motor (OpenCV/ONNX)
                    self.engine = get_engine()
                stage = "ocr"
                text, dropped = _read_embedded(data, self.name, page_like, self.engine)
                self.dropped += dropped
            except Exception as e:
                if isinstance(e, ConversionError) and e.code == "ocr_engine_unavailable":
                    raise  # sin motor no hay nada que leer: error del archivo, no de una imagen
                issue = PageIssue("image_failed", e.stage if isinstance(e, ConversionError) else stage, page=page)
                log_issue(issue, self.name, e)
                traceback.print_exc(file=sys.stderr)
                self.warnings.append(issue)
                text = ""
        self.done += 1
        self.on_progress(min(99, 100 * self.done // self.total))
        return text


# ---------------------------------------------------------------------------------------------------------------
# DOCX
# ---------------------------------------------------------------------------------------------------------------

def _docx_geometry(zf: zipfile.ZipFile) -> tuple[dict[str, float], float, int]:
    """(área mostrada en pulgadas² por SHA-1 de la imagen, área de página en pulgadas², nº de imágenes en el cuerpo).
    Si no se puede leer, ({}, Carta, 0): los llamadores caen al respaldo en píxeles."""
    import xml.etree.ElementTree as ET

    ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main",
          "wp": "http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing",
          "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
          "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships"}
    letter = 8.5 * 11.0
    try:
        root = ET.fromstring(zf.read("word/document.xml"))
        rels = ET.fromstring(zf.read("word/_rels/document.xml.rels"))
        targets = {r.get("Id"): r.get("Target") for r in rels}
        page_area = letter
        pg = root.findall(".//w:sectPr/w:pgSz", ns)
        if pg:
            w, h = pg[-1].get(f"{{{ns['w']}}}w"), pg[-1].get(f"{{{ns['w']}}}h")
            if w and h:
                page_area = (int(w) / 1440) * (int(h) / 1440)  # twips -> pulgadas
        areas: dict[str, float] = {}
        count = 0
        for drawing in list(root.iter(f"{{{ns['wp']}}}inline")) + list(root.iter(f"{{{ns['wp']}}}anchor")):
            extent = drawing.find("wp:extent", ns)
            blip = drawing.find(".//a:blip", ns)
            if extent is None or blip is None:
                continue
            target = targets.get(blip.get(f"{{{ns['r']}}}embed"))
            if not target:
                continue
            count += 1
            member = "word/" + target if not target.startswith("/") else target.lstrip("/")
            digest = hashlib.sha1(zf.read(member.replace("word/../", ""))).hexdigest()
            area = int(extent.get("cx")) * int(extent.get("cy")) / _EMU_PER_INCH ** 2
            areas[digest] = max(areas.get(digest, 0.0), area)
        return areas, page_area, count
    except Exception:
        return {}, letter, 0


def convert_docx(file_path: str, cancel_flag: threading.Event, on_progress: Callable[[int], None],
                 engine: OcrEngine | None = None) -> ConversionResult:
    path = Path(file_path)
    name = path.name
    if not path.exists():
        raise FileNotFoundError(f"Archivo no encontrado: {file_path}")
    if path.stat().st_size == 0:
        raise ValueError("El archivo está vacío")

    try:
        zf = zipfile.ZipFile(path)
    except zipfile.BadZipFile:
        return ConversionResult(convert_document(file_path))  # el mismo error de siempre
    with zf:
        media = [n for n in zf.namelist() if n.startswith("word/media/")]
        if not media:
            return ConversionResult(convert_document(file_path))  # ruta rápida: nada que reconocer
        areas, page_area, blips = _docx_geometry(zf)

    try:
        import mammoth
        from bs4 import BeautifulSoup
        from markitdown.converter_utils.docx.pre_process import pre_process_docx
        from markitdown.converters import HtmlConverter
    except Exception as e:  # dependencia interna de markitdown ausente: ruta estable sin OCR
        print(f"[Sidecar][ocr] file={name} stage=assemble code=office_fallback err={type(e).__name__}: {e}", file=sys.stderr)
        return ConversionResult(convert_document(file_path))

    reader = _Reader(name, cancel_flag, on_progress, engine, blips or len(media))
    results: dict[int, tuple[str, str]] = {}   # índice del marcador -> (alt, texto)

    def convert_image(image):
        with image.open() as f:
            data = f.read()
        digest = hashlib.sha1(data).hexdigest()
        area = areas.get(digest)
        if area is not None:
            ratio = area / page_area
            relevant, page_like = ratio >= OFFICE_MIN_AREA, ratio >= SCAN_COVERAGE
        else:  # tamaño mostrado desconocido: respaldo en píxeles
            size = _pixel_size(data)
            relevant = bool(size) and min(size) >= FALLBACK_MIN_SIDE_PX and size[0] * size[1] >= FALLBACK_MIN_PIXELS
            page_like = bool(size) and min(size) >= FALLBACK_PAGE_LIKE_PX
        text = reader.read(data, image.content_type, relevant, page_like, None)
        index = len(results)
        results[index] = (real_alt(image.alt_text), text)
        return {"src": f"markify-img:{index}", "alt": ""}

    try:
        with open(path, "rb") as f:
            html = mammoth.convert_to_html(pre_process_docx(f), convert_image=mammoth.images.img_element(convert_image)).value
    except _Cancelled:
        return ConversionResult("", reader.finish())
    except ConversionError:
        raise
    except Exception as e:  # la cadena propia falló antes de poder leer: ruta estable sin OCR
        print(f"[Sidecar][ocr] file={name} stage=assemble code=office_fallback err={type(e).__name__}: {e}", file=sys.stderr)
        return ConversionResult(convert_document(file_path))

    markdown = HtmlConverter().convert_string(_substitute_images(html, results, BeautifulSoup)).markdown
    if not markdown.strip():
        raise ConversionError("no_text_found", "assemble", name)
    return ConversionResult(markdown, reader.finish())


def _substitute_images(html: str, results: dict[int, tuple[str, str]], BeautifulSoup) -> str:
    """Cambia cada <img src="markify-img:N"> por su alt (en cursiva) y su texto, o la quita. Dentro de una celda de
    tabla todo va en línea para no romper la tabla Markdown."""
    from html import escape

    soup = BeautifulSoup(html, "html.parser")
    for img in soup.find_all("img"):
        src = img.get("src", "")
        if not src.startswith("markify-img:"):
            continue
        alt, text = results.get(int(src.split(":")[1]), ("", ""))
        paragraphs = [p for p in text.split("\n\n") if p.strip()]
        only_child = img.parent is not None and img.parent.name == "p" and \
            len([c for c in img.parent.contents if getattr(c, "name", None) or str(c).strip()]) == 1
        target = img.parent if only_child else img
        if not alt and not paragraphs:
            target.decompose()
            continue
        if only_child and img.find_parent(["td", "th"]) is None:   # bloque: un <p> por párrafo
            parts = ([f"<p><em>{escape(alt)}</em></p>"] if alt else []) + [f"<p>{escape(p)}</p>" for p in paragraphs]
        else:                                                     # en línea (celda o dentro de un párrafo)
            parts = ([f"<em>{escape(alt)}</em>"] if alt else []) + [escape(p) for p in paragraphs]
            parts = [" ".join(parts)]
        target.replace_with(*BeautifulSoup("".join(parts), "html.parser").contents)
    return str(soup)


# ---------------------------------------------------------------------------------------------------------------
# PPTX
# ---------------------------------------------------------------------------------------------------------------

def _iter_pictures(shapes):
    """(forma imagen, colección que la contiene), con recursión en grupos. Incluye marcadores con imagen."""
    from pptx.enum.shapes import MSO_SHAPE_TYPE as T

    for shape in list(shapes):
        if shape.shape_type == T.GROUP:
            yield from _iter_pictures(shape.shapes)
        elif shape.shape_type == T.PICTURE or (shape.shape_type == T.PLACEHOLDER and hasattr(shape, "image")):
            yield shape, shapes


def _pptx_content_type(shape) -> str | None:
    """Tipo de contenido de la parte del paquete (image/x-emf, image/svg+xml…). `shape.image.content_type` lo
    deduce con Pillow y falla con formatos que Pillow no conoce."""
    try:
        return shape.part.related_part(shape._element.blip_rId).content_type
    except Exception:
        try:
            return shape.image.content_type
        except Exception:
            return None


def _replace_with_text(shape, shapes, text: str) -> None:
    """Sustituye la forma por un cuadro de texto en las mismas coordenadas y posición z (o la quita si no hay texto)."""
    element = shape._element
    parent = element.getparent()
    if text:
        box = shapes.add_textbox(shape.left or 0, shape.top or 0, shape.width or 0, shape.height or 0)
        box.text_frame.text = _XML_ILLEGAL.sub("", text)
        parent.remove(box._element)
        parent.insert(parent.index(element), box._element)
    parent.remove(element)


def convert_pptx(file_path: str, cancel_flag: threading.Event, on_progress: Callable[[int], None],
                 engine: OcrEngine | None = None) -> ConversionResult:
    path = Path(file_path)
    name = path.name
    if not path.exists():
        raise FileNotFoundError(f"Archivo no encontrado: {file_path}")
    if path.stat().st_size == 0:
        raise ValueError("El archivo está vacío")

    try:
        with zipfile.ZipFile(path) as zf:
            has_media = any(n.startswith("ppt/media/") for n in zf.namelist())
    except zipfile.BadZipFile:
        return ConversionResult(convert_document(file_path))  # el mismo error de siempre
    if not has_media:
        return ConversionResult(convert_document(file_path))  # ruta rápida: nada que reconocer

    try:
        from io import BytesIO
        from pptx import Presentation

        from converter import convert_stream
        presentation = Presentation(str(path))
        pictures = [(number, shape, shapes) for number, slide in enumerate(presentation.slides, 1)
                    for shape, shapes in _iter_pictures(slide.shapes)]
    except Exception as e:  # no se pudo abrir con python-pptx: ruta estable sin OCR (mismo error que hoy si es inválido)
        print(f"[Sidecar][ocr] file={name} stage=open code=office_fallback err={type(e).__name__}: {e}", file=sys.stderr)
        return ConversionResult(convert_document(file_path))
    if not pictures:
        return ConversionResult(convert_document(file_path))

    slide_area = max(1, (presentation.slide_width or 1) * (presentation.slide_height or 1))
    reader = _Reader(name, cancel_flag, on_progress, engine, len(pictures))
    try:
        for number, shape, shapes in pictures:
            try:
                data, content_type = shape.image.blob, _pptx_content_type(shape)
            except Exception as e:  # relación rota o imagen ilegible dentro del PPTX
                issue = PageIssue("image_failed", "decode", page=number)
                log_issue(issue, name, e)
                reader.warnings.append(issue)
                _replace_with_text(shape, shapes, "")
                continue
            ratio = (shape.width or 0) * (shape.height or 0) / slide_area
            text = reader.read(data, content_type, ratio >= OFFICE_MIN_AREA, ratio >= SCAN_COVERAGE, number)
            try:
                descr = shape._element._nvXxPr.cNvPr.attrib.get("descr", "")
            except Exception:
                descr = ""
            _replace_with_text(shape, shapes, "\n".join(image_paragraphs(real_alt(descr, shape.name), text)))
    except _Cancelled:
        return ConversionResult("", reader.finish())

    buffer = BytesIO()
    presentation.save(buffer)
    buffer.seek(0)
    return ConversionResult(convert_stream(buffer, ".pptx"), reader.finish())
